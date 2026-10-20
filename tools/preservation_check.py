import re, sys
from pptx import Presentation
def deck_text(p, notes=True):
    out=[]
    for s in Presentation(p).slides:
        for sh in s.shapes:
            if sh.has_text_frame: out.append(sh.text_frame.text)
            if getattr(sh,'has_table',False) and sh.has_table:
                for r in sh.table.rows:
                    for c in r.cells: out.append(c.text)
            if sh.has_chart:
                ch=sh.chart
                for pl in ch.plots:
                    for se in pl.series: out.append(" ".join(str(v) for v in se.values))
                    out.append(" ".join(pl.categories))
        if notes and s.has_notes_slide: out.append(s.notes_slide.notes_text_frame.text)
    return "\n".join(out)
PAT=[r"\$\s?\d[\d.,]*\s?[MBK]?", r"\d[\d.,]*\s?%", r"\d+\s?bps", r"\bT-\d+", r"\bD\+\d+", r"\bDay[ -]\d+", r"\bA\d+\b",
     r"\b\d{1,2} (?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]* ?\d{0,4}", r"\b\d+(?:\.\d+)? ?(?:days|weeks|months|people|pp|points)\b",
     r"\b(?:Q[1-4]'?\d\d)\b", r"\b\d+\.\d+\b"]
def toks(t):
    t=t.replace("‑","-").replace("–","-")
    s=set()
    for p in PAT:
        for m in re.findall(p,t): s.add(re.sub(r"\s+"," ",m.strip().rstrip(".,")))
    return s
old=deck_text(sys.argv[1]); new=deck_text(sys.argv[2])
extra="".join(open(f).read() for f in sys.argv[3:])
norm=lambda x: x.replace(" ","").lower()
newn=norm(new+extra); 
miss=sorted(t for t in toks(old) if norm(t) not in newn)
print("v4 tokens",len(toks(old)),"missing in v5(+extra):",len(miss))
for m in miss: 
    i=old.replace("‑","-").find(m); print(f"  {m!r:22} ctx: {old[max(0,i-70):i+50]!r}")
