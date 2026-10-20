"""final_pass_upload.py: in-place tweaks to the candidate's uploaded deck (cohesion, depth, numbers). Keeps their layout."""
import sys, copy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData

SRC, DST = sys.argv[1], sys.argv[2]
prs = Presentation(SRC)
S = list(prs.slides)
MUTED = RGBColor(0x6B, 0x5B, 0x82)
FALLBACK = []

def paragraphs(slide):
    for sh in slide.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                yield sh, p
        if getattr(sh, "has_table", False) and sh.has_table:
            for r in sh.table.rows:
                for c in r.cells:
                    for p in c.text_frame.paragraphs:
                        yield sh, p

def rep(idx, old, new, count=1):
    """Replace old with new on slide idx (1-based); run-level when possible."""
    n = 0
    for sh, p in paragraphs(S[idx - 1]):
        full = "".join(r.text for r in p.runs)
        if old not in full:
            continue
        done = False
        for r in p.runs:
            if old in r.text:
                r.text = r.text.replace(old, new); done = True; break
        if not done:
            p.runs[0].text = full.replace(old, new)
            for r in p.runs[1:]:
                r.text = ""
            FALLBACK.append((idx, old[:40]))
        n += 1
    assert n == count, (idx, old[:60], n)


def rebuild(idx, startswith, parts):
    """Rebuild a paragraph's runs: parts = [(text, bold, size_pt_or_None)], run[0] is the format template."""
    for sh, para in paragraphs(S[idx - 1]):
        full = "".join(r.text for r in para.runs)
        if not full.startswith(startswith):
            continue
        tmpl = para.runs[0]._r
        for r in para.runs[1:]:
            r._r.getparent().remove(r._r)
        for i, (txt, bold, size) in enumerate(parts):
            el = tmpl if i == 0 else copy.deepcopy(tmpl)
            if i > 0:
                tmpl.getparent().append(el)
        runs = para.runs
        for r, (txt, bold, size) in zip(runs, parts):
            r.text = txt; r.font.bold = bold
            if size: r.font.size = Pt(size)
        return
    raise AssertionError((idx, startswith))

def run_sizes(idx, startswith):
    for sh, para in paragraphs(S[idx - 1]):
        if "".join(r.text for r in para.runs).startswith(startswith):
            return para.runs[0].font.size, para.runs[-1].font.size
    raise AssertionError((idx, startswith))

def footer(idx, text, y=6.62, h=0.5, size=9):
    s = S[idx - 1]
    box = s.shapes.add_textbox(Inches(0.67), Inches(y), Inches(11.5), Inches(h))
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.1); tf.margin_top = tf.margin_bottom = Inches(0.05)
    r = tf.paragraphs[0].add_run(); r.text = text
    r.font.size = Pt(size); r.font.name = "Arial"; r.font.color.rgb = MUTED

def notes(idx, text):
    S[idx - 1].notes_slide.notes_text_frame.text = text

# ---------------------------------------------------------------- slide 2: approach and success measures
rep(2, "120d algo approach", "120-day gated approach")
rep(2, "Standing up a Planning CoE that works with the Countries and GMs/CFOs: one Nubank", "CoS & CG and Global Finance / FP&A stand up a Planning CoE with the Countries and GMs: one Nubank")
rep(2, "Our team gathers inputs, analysis, recommendations; CEO / CFO decide", "Our team recommends; the CEO and CFO decide.")
rep(2, "Decks need to ships 7 days before the Board.", "The book ships 7 days before the Board, every quarter.")
rep(2, "to drive (1) projects getting prioritized via methodology, (2) board meetings having an evolving, living agenda with "
       "high-quality materials, (3) the HQ project is successfully coordinated, and (4) the team is seen as a value-adding "
       "growth facilitator instead of corporate bureaucracy.",
       "to drive (1) 100% of discretionary spend released against an approved case, with cycle time under 40 days; (2) a living "
       "Board agenda, with the book out 7 days before every Board; (3) HQ gates 1 and 2 cleared by mid-December, nothing signed "
       "between gates; (4) the team rated twice a year by the GMs as a growth facilitator, not corporate bureaucracy. Decisions "
       "requested today: the $25M FY27 cap and its release schedule, the Board calendar, and the HQ gates.")

for sh in S[1].shapes:
    if sh.shape_type == 1 and abs(sh.top - Inches(4.8)) < Inches(0.05):
        sh.top, sh.height = Inches(4.72), Inches(2.0)
    if sh.has_text_frame and sh.text_frame.text.startswith("How will we measure"):
        sh.top, sh.height = Inches(4.78), Inches(1.88)

# ---------------------------------------------------------------- slide 3: hub
rep(3, "CoS & GC team", "CoS & CG team")
rep(3, "Global Planning hub and market/vertical teams report to Global Finance, with a dotted line to Country CFOs.", "CoS & CG partners with Global Finance / FP&A to stand up the CoE; the hub and market/vertical teams report to Global Finance, with a dotted line to Country CFOs.")
rep(3, "% of relevant spend with an approved business case", "100% of discretionary spend with an approved business case")
rebuild(3, "KPIs:", [("KPIs", True, None), (": cycle time under 40 days, forecast accuracy, and 100% of discretionary spend with an approved case.", False, None)])
for sh in S[2].shapes:
    if abs(sh.top - Inches(1.75)) < Inches(0.02) and abs(sh.left - Inches(0.67)) < Inches(0.02) and sh.shape_type == 1:
        sh.top, sh.height = Inches(1.45), Inches(3.8)
    if sh.has_text_frame and sh.text_frame.text.startswith("THE CENTRALIZED"):
        sh.top, sh.height = Inches(1.52), Inches(3.68)

# ---------------------------------------------------------------- slide 6: schedule
rep(6, "go/non-go", "go/no-go")
rep(6, "decided on Q4 data;", "decided at the H1 close on H1 data, only from profit above plan, inside the efficiency-ratio "
       "guide in force and with NPL 90+ on plan (the CFO decides; Risk can block);")
rep(6, "deployed after Mteam validation", "deployed after MTeam validation")
hs, bs = run_sizes(6, "$15M")
hpt, bpt = (hs.pt if hs else None), (bs.pt if bs else None)
rebuild(6, "$25M", [("$25M", True, hpt), (" - Released to Business units on track to beat the FY27 plan, decided at the H1 close on H1 data: only "
    "from profit above plan, inside the efficiency-ratio guide in force, with NPL 90+ on plan (the CFO decides; Risk can block). Remedial spend "
    "for units behind plan goes to the MTeam, with CFO prioritization.", False, bpt)])
rebuild(6, "$5M", [("$10M", True, hpt), (" - Set aside for strategic imperatives, special projects, regulatory requirements or similar extraordinary "
    "escalation factors, and deployed after MTeam validation.", False, bpt)])
for sh in S[5].shapes:
    if sh.has_chart:
        ch = sh.chart
        plot = ch.plots[0]
        cats = list(plot.categories)
        cd = CategoryChartData()
        cd.categories = cats
        cd.add_series("base", (0, 50, 25, 15, 0))
        cd.add_series("value", (60, 10, 25, 10, 15))
        ch.replace_data(cd)
        assert [pt.data_label.text_frame.text for pt in plot.series[1].points] == ["$60M", "$5M", "$25M", "$10M", "$15M"]
        plot.series[1].points[1].data_label.text_frame.paragraphs[0].runs[0].text = "$10M"
# re-space the three right-hand cards so the longer stage-gate text fits; keep the footer clear
boxes = {}
for sh in S[5].shapes:
    if sh.has_text_frame and sh.text_frame.text.strip():
        for key in ("FUND NOW", "STAGE-GATE", "DEFER"):
            if sh.text_frame.text.startswith(key):
                boxes[key] = sh
rects = [sh for sh in S[5].shapes if sh.shape_type == 1 and not sh.text_frame.text.strip()]
def rect_for(box):
    return min(rects, key=lambda r: abs(r.top - box.top) + abs(r.left - box.left))
geom = {"FUND NOW": (1.8, 1.45), "STAGE-GATE": (3.35, 1.95), "DEFER": (5.4, 1.1)}
for key, (top, h) in geom.items():
    box = boxes[key]; r = rect_for(box)
    r.top, r.height = Inches(top), Inches(h)
    box.top, box.height = Inches(top + 0.12), Inches(h - 0.17)
footer(6, "The two releases ($25M) are the cap. The 2H checkpoint (up to $25M) and the reserve ($10M) sit outside the plan and open only "
          "on the conditions stated, so every dollar of the $60M has a path. The 70/30 split follows mature cohorts: ~$31 of monthly "
          "ARPAC at eight years against $17.1 for the group (Q1'26 cohort data).", y=6.6, h=0.5)

# ---------------------------------------------------------------- slide 7: ways of working
rep(7, "CoS & GC lead", "CoS & CG lead")
rep(7, "GC Team reports.", "The CoS & CG team reports.")

# ---------------------------------------------------------------- slide 8: Board prep
rep(8, "Today I issue the locked outline", "We issue the locked outline")
rep(8, "The capital options memo settles the ranking and the 70/30 split with the CEO and CFO at T-16.",
       "The options memo settles the capital priorities and the 70/30 core / new-bets split with the CEO and CFO at T-16.")
rep(8, "Board assumed on Tue 17 Nov, four weeks from starting date. Q3'26 results should be due 12 Nov.",
       "Board assumed on Tue 17 Nov, four weeks from today. Q3'26 results are due 12 Nov; Investor Day is 8 Dec 2026.")

# ---------------------------------------------------------------- slide 9: friction
rep(9, "decision-making process better even . GMs", "decision-making process better. GMs")

# ---------------------------------------------------------------- slide 10: HQ assessment
rep(10, "so the purpose survives Pillar Two", "so the incentive survives Pillar Two")
rebuild(10, "4. Substance", [("4. ", False, None), ("Substance", True, None), (": real payroll and assets, so the incentive survives Pillar Two (Tax, People).", False, None)])
rep(13, " Review commitments and initial boundaries in speaker notes.", "")
footer(10, "Nu Holdings is Cayman-incorporated (Form 20-F). Nu is above the EUR 750M Pillar Two threshold; 15% minimum; the 2026 "
           "carve-out is 9.4% of payroll and 7.4% of tangible assets (OECD GloBE, Art. 5.3 and 9.2). Brazil applies a domestic "
           "minimum top-up tax from 2025. Jurisdiction open; no savings figure until a site is chosen.", y=6.93, h=0.45, size=8.5)

# ---------------------------------------------------------------- slide 11: officers and leadership
rep(11, "Decisions are made at the SteerCo level.", "Gate decisions are taken at the monthly SteerCo, chaired by the CFO.")
rep(11, "Unresolved issues reach the CFO within 48 hours. Legal and Compliance retain legality vetoes.",
        "When functions disagree, both views go in the log and reach the CFO within 48 hours with a recommendation; "
        "Legal and Compliance keep legality vetoes.")
for sh in S[10].shapes:
    if sh.has_text_frame and sh.text_frame.text.startswith("CROSS-FUNCTIONAL"):
        sh.height = Inches(3.42)
rep(11, "CoS & CG leads cadence and delivery. Functions own specialist / technical advice.", "CoS & CG leads cadence and delivery; functions own the technical advice.")
rep(11, "Weekly functional leads meeting, supported by one decision log; monthly SteerCo meetings, with senior-level engagement.",
        "Weekly functional leads meeting with one decision log; monthly SteerCo.")

# ---------------------------------------------------------------- slide 14
rep(14, "Geographies abd Functions", "Geographies and Functions")

# ---------------------------------------------------------------- slide 15: divider list
for sh in S[14].shapes:
    if sh.has_text_frame and "Full assumptions" in sh.text_frame.text:
        ps = sh.text_frame.paragraphs
        ps[0].runs[0].text = "A  Key assumptions   |   B  Statutory entity model   |   D  Sources"
        for r in ps[0].runs[1:]: r.text = ""
        ps[1].runs[0].text = "F  Day-by-day timeline   |   H  AI support   |   J  HQ site scorecard and risk register   |   K  How AI was used"
        for r in ps[1].runs[1:]: r.text = ""
        for p in ps[2:]:
            for r in p.runs: r.text = ""

# ---------------------------------------------------------------- backup A (16)
rep(16, "inside the ~20% guide, NPL 90+ on plan", "inside the efficiency-ratio guide in force, NPL 90+ on plan")
rep(16, "Option scores and weights are my judgment", "Option scores and weights are our judgment")
rep(16, "Slide 4; backups C, D", "Slide 5; backup D")
rep(16, "FY27 envelope: first release 1 Jan 2027, second at Q1 close (31 Mar), reserve at H1 close. Hub of 4-6 people.",
        "FY27 envelope: first release 1 Jan 2027, second at Q1 close (decided on Q1 data), 2H checkpoint (up to $25M) and reserve ($10M) at H1 close. Hub of 4-6 people.")
rep(16, "Calendar fiscal year; sizing is mine", "Calendar fiscal year; sizing is ours")
rep(16, "the site scorecard weights are my judgment (backup J)", "the site scorecard weights are our judgment (backup J)")

# ---------------------------------------------------------------- backup D (18)
rep(18, "Slides 3, 4, 7; backup I", "Slides 5, 7; backup A")
rep(18, "Slides 7, 9, 15; backup B", "Slides 9, 13; backup B")
rep(18, "Slide 7; backup E", "Slide 7 (70/30 split)")
rep(18, "Slides 6, 9, 12; backup H", "Slides 2, 3, 8, 14; backup H")
rep(18, "Slide 10", "Slide 11")
rep(18, "Slide 5", "Slide 6")

# ---------------------------------------------------------------- backup F (19)
rep(19, "Hub in place; finance partners embedded", "Hub in place; market and vertical finance teams aligned")
rep(19, "Squad clinics running; $12M case track under way", "Bi-weekly stand-ups running; business cases in review")
rep(19, "Gate 2: grant terms (24-month grace); one site dropped", "Gate 2: grant terms (12-24 months' grace); one site dropped")

# ---------------------------------------------------------------- backup J (21)
rep(21, "Proportional scaling and 24 months' grace", "Proportional scaling and 12-24 months' grace")
rep(21, "Dependence on me", "Dependence on the CoS")
rep(21, "gate packs the CFO can read without me", "gate packs the CFO can decide on without the lead")
rep(21, "Me", "CoS & CG")
rep(21, "Weights and mitigations are my proposals (A13).", "Weights and mitigations are our proposals.")


# ---------------------------------------------------------------- new slide after the hub: structure, leveling, hiring waves
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
PURPLE, INK, BODY, PAGE, BAND, WHITE, LINE = (RGBColor(0x82,0x0A,0xD1), RGBColor(0x11,0x05,0x2C), RGBColor(0x3B,0x2A,0x54),
    RGBColor(0xFA,0xF8,0xFF), RGBColor(0xF5,0xF0,0xFC), RGBColor(0xFF,0xFF,0xFF), RGBColor(0xE2,0xD9,0xF3))
blank = [l for l in prs.slide_layouts if l.name == "Blank"][0]
ns = prs.slides.add_slide(blank)
for shp in list(ns.shapes):
    shp._element.getparent().remove(shp._element)
ns.background.fill.solid(); ns.background.fill.fore_color.rgb = PAGE
def tx(x, y, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    b = ns.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = b.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.1); tf.margin_top = tf.margin_bottom = Inches(0.05); tf.vertical_anchor = anchor
    runs = text if isinstance(text, list) else [(text, bold)]
    para = tf.paragraphs[0]; para.alignment = align
    for s_, b_ in runs:
        r = para.add_run(); r.text = s_; r.font.size = Pt(size); r.font.bold = b_; r.font.name = "Arial"; r.font.color.rgb = color
    return b
tx(0.67, 0.38, 12.0, 0.3, "DELIVERABLE 1  //  TEAM", 10, PURPLE, bold=True)
tx(0.67, 0.68, 12.0, 0.95, "Hub of 4-6: 2-3 in the countries, 2-3 central in the US, in two waves", 24, INK, bold=True)
rows = [["Role", "Level (proposal)", "Where", "Reports to", "Wave"],
        ["Hub lead: CoS & Corporate Governance", "Existing role", "US (central)", "Global CFO", "In place"],
        ["Global planning manager: owns the model, the case template and the release packs", "Senior manager", "US (central)", "Hub lead", "1: by Q2'27"],
        ["Finance partner, Brazil: cases and drivers with the GMs and the Brazil CFO", "Manager / senior manager", "Brazil", "Global Finance; dotted line to the Country CFO", "1: by Q2'27"],
        ["Finance partner, Mexico: same role for Mexico (and Colombia until wave 2)", "Manager", "Mexico", "Global Finance; dotted line to the Country CFO", "1: by Q2'27"],
        ["Decision-support analyst: data platform, AI agents, consolidation", "Senior analyst", "US (central)", "Global planning manager", "2: by Q4'27, if gates pass"],
        ["Finance partner, Colombia", "Manager", "Colombia", "Global Finance; dotted line to the Country CFO", "2: by Q4'27, if gates pass"],
        ["Second central analyst (only if volume requires)", "Analyst", "US (central)", "Global planning manager", "2: by Q4'27, if gates pass"]]
col_w = [4.6, 1.9, 1.3, 2.6, 1.7]
shp = ns.shapes.add_table(len(rows), len(col_w), Inches(0.67), Inches(1.7), Inches(sum(col_w)), Inches(0.38 * len(rows)))
tbl = shp.table; tbl.first_row = False; tbl.horz_banding = False
tblPr = tbl._tbl.tblPr
for c, wdt in enumerate(col_w):
    tbl.columns[c].width = Inches(wdt)
for r_i, row in enumerate(rows):
    tbl.rows[r_i].height = Inches(0.42 if r_i else 0.36)
    for c_i, val in enumerate(row):
        cell = tbl.cell(r_i, c_i); cell.fill.solid()
        cell.fill.fore_color.rgb = PURPLE if r_i == 0 else (WHITE if r_i % 2 else BAND)
        cell.margin_left = cell.margin_right = Inches(0.08); cell.margin_top = cell.margin_bottom = Inches(0.04)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        para = cell.text_frame.paragraphs[0]; run = para.add_run(); run.text = val
        run.font.name = "Arial"; run.font.size = Pt(10.5 if r_i else 11)
        run.font.bold = (r_i == 0) or (c_i == 0 and r_i > 0)
        run.font.color.rgb = WHITE if r_i == 0 else (INK if c_i == 0 else BODY)
band = ns.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.67), Inches(5.25), Inches(11.8), Inches(1.15))
band.fill.solid(); band.fill.fore_color.rgb = BAND; band.line.color.rgb = LINE; band.line.width = Pt(0.75)
band.adjustments[0] = 0.08; band.shadow.inherit = False; band.text_frame.text = ""
tx(0.87, 5.35, 11.4, 0.98, [("Quality gates for wave 2 (checked at the H1 close): ", True),
   ("wave 1 in seat and onboarded; cycle time under 40 days on the Q1'27 release; 100% of first-release spend with an approved "
    "case; forecast-accuracy baseline set in Q1 and improving in Q2; mid-year GM feedback positive. If a gate fails, wave 2 waits "
    "and the gap is fixed first.", False)], 11.5, INK, anchor=MSO_ANCHOR.MIDDLE)
tx(0.67, 6.55, 11.5, 0.45, "Proposal. CoS & CG partners with Global Finance / FP&A to stand up the CoE; the central roles sit with Global Finance in the US. "
   "Levels map to Nu's leveling at hire. Country partners are market hires, not relocations. Hiring cost is central opex and counts against the "
   "$25M cap. Wave 1 recruiting starts after the Board; wave 2 only on the gates above.", 9, MUTED)
# page number box, copied from slide 3's
src_pn = [sh for sh in S[2].shapes if sh.has_text_frame and sh.text_frame.text.strip().isdigit()][0]
ns.shapes._spTree.append(copy.deepcopy(src_pn._element))
# move the new slide to position 4
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)
new_id = ids[-1]; sldIdLst.remove(new_id); sldIdLst.insert(3, new_id)
S = list(prs.slides)


# ---------------------------------------------------------------- last backup K: how AI was used
ns = prs.slides.add_slide(blank)
for shp in list(ns.shapes):
    shp._element.getparent().remove(shp._element)
ns.background.fill.solid(); ns.background.fill.fore_color.rgb = PAGE
tx(0.67, 0.38, 12.0, 0.3, "BACKUP K", 10, PURPLE, bold=True)
tx(0.67, 0.68, 12.0, 0.95, "How AI was used: an AI-supported case, not an AI-driven one", 24, INK, bold=True)
def card_k(x, label, heading, lines):
    r = ns.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.75), Inches(5.8), Inches(3.35))
    r.fill.solid(); r.fill.fore_color.rgb = WHITE; r.line.color.rgb = LINE; r.line.width = Pt(0.75)
    r.adjustments[0] = 0.05; r.shadow.inherit = False; r.text_frame.text = ""
    b = ns.shapes.add_textbox(Inches(x + 0.12), Inches(1.87), Inches(5.56), Inches(3.1)); tf = b.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.1); tf.margin_top = tf.margin_bottom = Inches(0.05)
    specs = [(label, 9.5, PURPLE, True, 3), (heading, 15, INK, True, 6)] + [(l, 12, BODY, False, 4) for l in lines]
    for i, (txt, size, color, bold, after) in enumerate(specs):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); para.space_after = Pt(after)
        run = para.add_run(); run.text = txt; run.font.size = Pt(size); run.font.bold = bold; run.font.name = "Arial"; run.font.color.rgb = color
card_k(0.67, "WHAT AI SUPPORTED", "Research, checking and tidying",
       ["Research on Nu's public data, filings and announcements, and on planning and governance practices.",
        "Numbers: calculations re-run and cross-checked against the sources.",
        "References and best practices: sources found, read and summarized.",
        "Minor formatting of slides and tables.",
        "Testing some of the assumptions and scenarios (for example the option-weight sensitivity)."])
card_k(6.67, "WHAT WAS HUMAN-BUILT", "The thinking, the writing and the calls",
       ["The ideas, recommendations and judgments, and all the writing.",
        "The structure, the storyline and the framing of each deliverable.",
        "The assumptions, their classification as fact or assumption, and the trade-offs taken.",
        "The final review of every figure, reference and recommendation.",
        "The decision to present this as an AI-supported case rather than let AI drive it."])
band = ns.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.67), Inches(5.3), Inches(11.8), Inches(1.2))
band.fill.solid(); band.fill.fore_color.rgb = BAND; band.line.color.rgb = LINE; band.line.width = Pt(0.75)
band.adjustments[0] = 0.08; band.shadow.inherit = False; band.text_frame.text = ""
tx(0.87, 5.4, 11.4, 1.0, [("Tools: ", True), ("Claude, Gemini and Copilot were used, each for different aspects of the work; no single tool "
   "produced the case. ", False), ("Looking ahead: ", True), ("the AI agents proposed for the planning hub and the Board and HQ work (slide 3, "
   "backup H) can be set up in any LLM Nubank enables, inside Nu's environment and on approved tools, with people owning every "
   "number and decision.", False)], 11.5, INK, anchor=MSO_ANCHOR.MIDDLE)
tx(0.67, 6.62, 11.5, 0.4, "Every figure, calculation, reference and recommendation in this deck was reviewed and is owned by the author.", 9, MUTED)
ns.shapes._spTree.append(copy.deepcopy(src_pn._element))
S = list(prs.slides)

# ---------------------------------------------------------------- page numbers
for i, s in enumerate(S, 1):
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip().isdigit():
            sh.text_frame.paragraphs[0].runs[0].text = str(i)

# ---------------------------------------------------------------- no speaker notes anywhere (slides only)
for s in S:
    if s.has_notes_slide:
        s.notes_slide.notes_text_frame.clear()
        s.notes_slide.notes_text_frame.text = ""
prs.save(DST)
print("saved", DST, "slides", len(prs.slides), "fallback paragraphs:", FALLBACK)
