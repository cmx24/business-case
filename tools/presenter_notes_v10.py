"""presenter_notes_v10.py: copy of the v10 deck with short presenter cue cards in the notes (slides 1-15). The send-out deck stays notes-free."""
import sys
from pptx import Presentation
SRC, DST = sys.argv[1], sys.argv[2]
C = {
1: ("1.0 min", "Three problems at once; finance and governance as growth levers.",
    ["$60M of asks vs a $25M cap; a Board in 4 weeks with two executives disagreeing; a new HQ with no line authority.",
     "Add the rigor scale needs without losing Nu's speed."],
    "Reading the three cards.", "Here's how we'd approach all four."),
2: ("1.6 min", "One principle, ownership, used four times; three decisions today.",
    ["Options with a recommendation; decision, reasons and dissent written down.",
     "Measures: 100% of spend with a case; book out 7 days early; HQ gates 1-2 by mid-Dec; GMs rate us twice a year.",
     "Ask: $25M cap and release schedule, Board calendar, HQ gates."],
    "Explaining every tile.", "Deliverable 1 of 4: planning."),
3: ("1.6 min", "A small hub, built with Global Finance / FP&A; markets own their numbers.",
    ["4-6 people; report to Global Finance, dotted line to country CFOs.",
     "KPIs: cycle time < 40 days, forecast accuracy, 100% of spend with a case.",
     "Three meetings, nothing else on the calendar. Local finance owns drivers; we consolidate; CFO decides."],
    "'My team'. Say 'the hub'.", "Who the team is, and when we hire."),
4: ("1.2 min", "Hire half now, half only if the quality gates pass.",
    ["Wave 1 by Q2'27: planning manager (US), Brazil and Mexico partners.",
     "Wave 2 by Q4'27 if gates pass: analyst (US), Colombia, optional 2nd analyst.",
     "Gates: <40-day cycle, 100% with a case, accuracy improving, positive GM feedback. Cost counts against the $25M."],
    "Defending each role.", "Now the money: the $35M gap."),
5: ("1.8 min", "$35M is small for Nu; the precedent matters, so budgets work like Nu's credit.",
    ["~15 bps of revenue; ~0.2 pts of efficiency ratio. Say it calmly.",
     "Broken today: asks arrive as one number, no shared format.",
     "Equal cut teaches padding; core-first starves new bets; staged release recommended. 20-F low-and-grow."],
    "Overselling the $35M.", "How a case gets scored and decided."),
6: ("1.6 min", "Mandatory spend has its own gate; discretionary asks compete on published criteria.",
    ["Five criteria; economic value highest at 30%. Published before submission.",
     "The score informs; it doesn't decide.",
     "Submit -> calibrate -> CFO approves within cap, CEO settles trade-offs -> track monthly, reallocate quarterly."],
    "Walking through the formula.", "How the money flows through the year."),
7: ("2.0 min", "$25M is the cap in two releases; the rest has a conditional path.",
    ["1Q27 $15M (70/30 core/new bets); 2Q27 $10M on Q1 data = the $25M cap.",
     "2H checkpoint up to $25M: only from profit above plan, inside the guide, NPL on plan; CFO decides, Risk can block.",
     "$10M reserve after MTeam validation. Nobody hears a flat no."],
    "Implying the full $60M is approved.", "PAUSE for questions. Then Deliverable 2 of 4."),
8: ("1.4 min", "Clear roles, four rhythms; executives own the content.",
    ["Lead brings options and a recommendation; core team runs calendar, briefs, actions.",
     "Chair approves the agenda; Corporate Secretary owns formal governance.",
     "Weekly CFO; monthly CFO + GMs; fixed pre-Board countdown; actions within 5 working days after."],
    "'I resolve executive differences'.", "Here's the countdown for this Board."),
9: ("1.8 min", "The book is late because nothing is locked; a 28-day countdown fixes it, every quarter.",
    ["28 days out: outline locked. 16: executive lock. 10: inputs freeze. 7: book on the portal.",
     "Can't slip: the capital plan (last Board before Investor Day, 8 Dec) and statutory items.",
     "Still no agreement at the lock: both positions in the memo, CEO decides, dissent recorded."],
    "Saying 'T-minus'. Say 'days out'.", "The people side of disagreement."),
10: ("1.8 min", "Rules public, hard conversations private, sign-off joint.",
    ["Talk privately first. Move past the number to the evidence.",
     "People: access to teams, not just leaders. Build with them, not for them.",
     "[Your example, ~40 words, if ready.] Disagree and commit is a good outcome."],
    "'I win them over'.", "PAUSE for questions. Then Deliverable 3 of 4."),
11: ("1.7 min", "Governance first, campus last; hard tests before preferences.",
    ["Two kinds of agreement: regulatory and campus.",
     "Four pass/fail tests: licenses, tax residence, officers, real substance.",
     "Pillar Two: only ~9% of payroll and ~7% of assets carved out in 2026. Two sites to term sheet; CEO and CFO choose."],
    "Naming a city.", "Who owns which call."),
12: ("1.5 min", "Experts own the judgment, the team owns the process, the CFO owns the recommendation.",
    ["Monthly SteerCo chaired by the CFO; signing central until officers approved and insured.",
     "Filings start at gate 1: approvals are the long pole.",
     "Disagreement: both views to the CFO within 48h. Legal and Compliance veto on legality only."],
    "Implying you'd sign or be an officer.", "And the four gates."),
13: ("1.7 min", "Four decisions, each with an owner, an exit test and a walk-away.",
    ["Can we? (Tax, day 30). What may we offer? (Legal, day 60). Who may bind Nu? (day 90). Do we commit? (lease, day 120).",
     "One walk-away: any site that moves tax residence is out, whatever the incentive.",
     "Nothing signed between gates; decisions stay with the MTeam."],
    "Reading every cell.", "Deliverable 4 of 4: how this function works."),
14: ("1.7 min", "How the function works: trust, clear recommendations, visible accountability, learning.",
    ["Criteria explicit before the decision; trade-offs to the CFO with a recommendation.",
     "Every commitment has an owner and a date; AI supports, people validate.",
     "[Your example, ~40 words: when the data showed you were wrong.] Local decision rights stay where they are."],
    "Personal traits ('I'm collaborative').", "That in practice, and how we'll know it's working."),
15: ("1.6 min", "The function earns its place with the markets; then the ask.",
    ["Team owns the process; GMs and local CFOs own their numbers and decisions.",
     "Quarterly with the CFO: keep, fix or drop.",
     "ASK: $25M cap in two releases, the Board calendar, the HQ gates. Report back in 90 days."],
    "Ending on 'thank you' alone.", "Go back to slide 2 for Q&A."),
}
prs = Presentation(SRC)
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

import copy
TEMPLATE = None  # a notes page that has the slide-image and body placeholders

def notes_frame(slide):
    """Return the notes text frame; if the notes page has no body placeholder (the file's notes master has none),
    copy the slide-image and body placeholders from a complete notes page."""
    ns = slide.notes_slide
    if ns.notes_text_frame is not None:
        return ns.notes_text_frame
    for sh in TEMPLATE.shapes:
        if sh.is_placeholder and sh.placeholder_format.type in (101, 2):
            ns.shapes._spTree.append(copy.deepcopy(sh._element))
    tf = ns.notes_text_frame
    tf.text = ""
    return tf

for s in prs.slides:
    if s.has_notes_slide and s.notes_slide.notes_text_frame is not None and len(s.notes_slide.placeholders) >= 2:
        TEMPLATE = s.notes_slide
        break
for i, s in enumerate(prs.slides, 1):
    tf = notes_frame(s)
    if i in C:
        time, land, says, avoid, nxt = C[i]
        lines = [f"TIME {time}  |  Say 'we', not 'I'", f"LAND: {land}"] + [f"- {x}" for x in says] + [f"AVOID: {avoid}", f"NEXT: {nxt}"]
        tf.text = "\n".join(lines)
    else:
        tf.text = ""
prs.save(DST)
print("saved", DST)
