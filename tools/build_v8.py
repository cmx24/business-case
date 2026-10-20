#!/usr/bin/env python3
"""build_v8.py - builds deck/CFO_CoS_deck_v8.pptx: v7 reframed as one way of deciding used three times (Fable reassessment and final pass).

Uses deck v4 only as a template (theme, master, notes master); every slide is drawn here, so all wording is in this file.
A font-metric fit check (Liberation Sans = Arial metrics) warns when any text box or table would overflow.
usage: python3 -I tools/build_v8.py [dst]
"""
import os
import sys

from PIL import ImageFont
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import pptx_safe  # noqa: E402

TEMPLATE = "deck/CFO_CoS_deck_v4.pptx"
DST = sys.argv[1] if len(sys.argv) > 1 else "deck/CFO_CoS_deck_v8.pptx"

DARK = RGBColor(0x19, 0x05, 0x30)
PURPLE = RGBColor(0x82, 0x0A, 0xD1)
LILAC_TXT = RGBColor(0xC0, 0x84, 0xFC)
INK = RGBColor(0x11, 0x05, 0x2C)
BODY = RGBColor(0x3B, 0x2A, 0x54)
MUTED = RGBColor(0x6B, 0x5B, 0x82)
LINE = RGBColor(0xE2, 0xD9, 0xF3)
PAGE = RGBColor(0xFA, 0xF8, 0xFF)
BAND = RGBColor(0xF7, 0xF2, 0xFD)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CARD_DARK = RGBColor(0x2A, 0x0C, 0x4E)
ORANGE = RGBColor(0xD9, 0x77, 0x06)
TEAL = RGBColor(0x0F, 0x76, 0x6E)
SOFT = RGBColor(0xE9, 0xD5, 0xFF)

FONT = "Arial"
REG = ImageFont.truetype("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf", 100)
BLD = ImageFont.truetype("/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf", 100)
WARN = []

prs = Presentation(TEMPLATE)
for sldId in list(prs.slides._sldIdLst):
    prs.part.drop_rel(sldId.rId)
    prs.slides._sldIdLst.remove(sldId)
BLANK = [lay for lay in prs.slide_layouts if lay.name == "Blank"][0]
CUR = {"n": 0, "title": ""}


# ------------------------------------------------------------------ measurement
def text_width_pt(text, size, bold):
    f = BLD if bold else REG
    return f.getlength(text) * size / 100.0


def wrapped_lines(runs, size, width_pt):
    """runs: list of (text, bold). Greedy word wrap with per-run fonts; returns line count."""
    words = []
    for text, bold in runs:
        for i, w in enumerate(text.split(" ")):
            words.append((w, bold, i > 0 or (words and text.startswith(" "))))
    lines, cur = 1, 0.0
    space = text_width_pt(" ", size, False)
    for w, bold, _ in words:
        if not w:
            continue
        ww = text_width_pt(w, size, bold)
        add = ww if cur == 0 else space + ww
        if cur + add > width_pt and cur > 0:
            lines += 1
            cur = ww
        else:
            cur += add
    return lines


def para_runs(p):
    if isinstance(p, str):
        return [(p, False)]
    return [(r[0], r[1]) for r in p]


def check_fit(where, paras, sizes, w_in, h_in, margin_lr=0.1, margin_tb=0.05, spacing=1.2, space_after=None):
    width_pt = (w_in - 2 * margin_lr) * 72
    total = 0.0
    for i, p in enumerate(paras):
        size = sizes[i]
        runs = para_runs(p)
        n = wrapped_lines(runs, size, width_pt)
        total += n * size * spacing
        if space_after:
            total += space_after[i]
    need = total / 72 + 2 * margin_tb
    if need > h_in + 0.02:
        WARN.append(f"slide {CUR['n']} '{CUR['title'][:40]}' {where}: needs {need:.2f}in, box {h_in:.2f}in")
    return need


# ------------------------------------------------------------------ primitives
def new_slide(dark=False, title=""):
    s = prs.slides.add_slide(BLANK)
    for shp in list(s.shapes):
        shp._element.getparent().remove(shp._element)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = DARK if dark else PAGE
    CUR["n"] = len(prs.slides)
    CUR["title"] = title
    return s


def add_text(s, x, y, w, h, paras, size=12.5, color=BODY, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             sizes=None, colors=None, space_after=None, where="text", check=True):
    """paras: list of str or list of (text, bold[, color]) runs."""
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.1)
    tf.margin_top = tf.margin_bottom = Inches(0.05)
    tf.vertical_anchor = anchor
    sizes = sizes or [size] * len(paras)
    colors = colors or [color] * len(paras)
    sa = space_after or [0] * len(paras)
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if sa[i]:
            p.space_after = Pt(sa[i])
        runs = [(para, bold)] if isinstance(para, str) else para
        for run in runs:
            r = p.add_run()
            r.text = run[0]
            r.font.size = Pt(sizes[i])
            r.font.bold = bool(run[1])
            r.font.name = FONT
            r.font.color.rgb = run[2] if len(run) > 2 else colors[i]
    if check:
        conv = [[(r[0], r[1]) for r in ([(p, bold)] if isinstance(p, str) else p)] for p in paras]
        check_fit(where, conv, sizes, w, h, space_after=sa)
    return box


def rect(s, x, y, w, h, fill=WHITE, line=LINE, rounded=True, adj=0.05):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                             Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(0.75)
    if rounded:
        shp.adjustments[0] = adj
    shp.shadow.inherit = False
    shp.text_frame.text = ""
    return shp


def header(s, label, title, dark=False):
    add_text(s, 0.67, 0.38, 12.0, 0.3, [label], size=10, bold=True, color=(LILAC_TXT if dark else PURPLE), where="label")
    add_text(s, 0.67, 0.68, 12.0, 0.95, [title], size=24, bold=True, color=(WHITE if dark else INK), where="title")
    n = wrapped_lines([(title, True)], 24, (12.0 - 0.2) * 72)
    CUR.setdefault("title_bottom", {})[CUR["n"]] = 0.68 + 0.05 + n * 24 * 1.2 / 72


def footer(s, text, y=6.88, h=0.42):
    add_text(s, 0.67, y, 11.5, h, [text], size=9, color=MUTED, where="footer")


def card(s, x, y, w, h, label, heading, body, label_color=PURPLE, body_size=12, heading_size=15, fill=WHITE, line=LINE,
         text_color=BODY, heading_color=INK):
    rect(s, x, y, w, h, fill=fill, line=line)
    paras, sizes, colors, sa = [], [], [], []
    if label:
        paras.append([(label, True)]); sizes.append(9.5); colors.append(label_color); sa.append(3)
    if heading:
        paras.append([(heading, True)]); sizes.append(heading_size); colors.append(heading_color); sa.append(6)
    for b in body:
        paras.append(b); sizes.append(body_size); colors.append(text_color); sa.append(4)
    add_text(s, x + 0.12, y + 0.12, w - 0.24, h - 0.2, paras, sizes=sizes, colors=colors, space_after=sa, where=f"card '{heading or label}'")


def band(s, x, y, w, h, paras, size=12.5, dark=True):
    rect(s, x, y, w, h, fill=(DARK if dark else BAND), line=(None if dark else LINE), adj=0.08)
    add_text(s, x + 0.2, y + 0.1, w - 0.4, h - 0.16, paras, size=size, color=(WHITE if dark else INK), anchor=MSO_ANCHOR.MIDDLE,
             where="band")


def table(s, x, y, col_w, rows, header_row=True, size=10.5, header_size=11, first_bold=True, max_bottom=6.8,
          highlight_last=False, row_min=0.3):
    nrows, ncols = len(rows), len(rows[0])
    shp = s.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(sum(col_w)), Inches(row_min * nrows))
    tbl = shp.table
    tbl.first_row = False
    tbl.horz_banding = False
    for c, wdt in enumerate(col_w):
        tbl.columns[c].width = Inches(wdt)
    total_h = 0.0
    for r, row in enumerate(rows):
        row_h = 0.0
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.07)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            is_head = header_row and r == 0
            cell.fill.solid()
            if is_head:
                cell.fill.fore_color.rgb = PURPLE
            elif highlight_last and r == nrows - 1:
                cell.fill.fore_color.rgb = SOFT
            else:
                cell.fill.fore_color.rgb = WHITE if r % 2 else BAND
            tf = cell.text_frame
            tf.word_wrap = True
            runs = [(val, is_head or (first_bold and c == 0))] if isinstance(val, str) else val
            p = tf.paragraphs[0]
            for t, b in runs:
                rr = p.add_run()
                rr.text = t
                rr.font.size = Pt(header_size if is_head else size)
                rr.font.bold = bool(b)
                rr.font.name = FONT
                rr.font.color.rgb = WHITE if is_head else (INK if c == 0 else BODY)
            sz = header_size if is_head else size
            n = wrapped_lines([(t, b) for t, b in runs], sz, (col_w[c] - 0.14) * 72)
            row_h = max(row_h, n * sz * 1.2 / 72 + 0.08)
        row_h = max(row_h, row_min)
        tbl.rows[r].height = Inches(row_h)
        total_h += row_h
    shp.height = Inches(total_h)
    if y + total_h > max_bottom + 0.02:
        WARN.append(f"slide {CUR['n']} '{CUR['title'][:40]}' table bottom {y + total_h:.2f}in > {max_bottom}in")
    return y + total_h


def notes(s, text):
    s.notes_slide.notes_text_frame.text = text


def number_all():
    for i, sl in enumerate(prs.slides, 1):
        dark = sl.background.fill.fore_color.rgb == DARK
        add_text(sl, 12.35, 7.05, 0.8, 0.3, [str(i)], size=9, color=(LILAC_TXT if dark else MUTED), align=PP_ALIGN.RIGHT,
                 check=False)


# ================================================================== 1. Title
s = new_slide(dark=True, title="title")
add_text(s, 0.67, 0.75, 12.0, 0.3, ["NU HOLDINGS  //  OFFICE OF THE CFO AND CORPORATE GOVERNANCE"], size=10.5, bold=True,
         color=LILAC_TXT, where="label")
add_text(s, 0.67, 1.1, 11.5, 1.3, ["Chief of Staff to the CFO: planning, the Board and the new HQ"], size=36, bold=True,
         color=WHITE, where="title")
add_text(s, 0.67, 2.45, 10.5, 0.75, ["Case presentation, 60 days into the role. Cassio Modelli, October 2026. Nu figures are "
                                     "public Q2'26 data; my assumptions are on slide 3."], size=15, color=SOFT, where="sub")
cards1 = [("PLANNING", "$60M of asks, a $25M cap", "Bring planning in-house and give every ask an answer."),
          ("GOVERNANCE", "A Board in four weeks", "Two executives disagree on capital, and the Board book is late."),
          ("SPECIAL PROJECT", "A new HQ", "Legal, tax, compliance and campus terms across jurisdictions, with no line authority.")]
for i, (lab, head, body) in enumerate(cards1):
    card(s, 0.67 + i * 4.05, 3.45, 3.8, 1.75, lab, head, [body], label_color=LILAC_TXT, fill=CARD_DARK, line=None,
         text_color=SOFT, heading_color=WHITE, body_size=13)
band(s, 0.67, 5.45, 11.9, 0.95, [[("How I triaged them: ", True), ("the Board date is fixed, so it goes first. The budget "
     "sets the precedent for every cycle after it. The HQ is chartered now and paced by gates.", False)]], size=13.5,
     dark=False)
notes(s, "Triage: the Board date is fixed, so it goes first; the budget sets the precedent; the HQ is paced by gates.")

# ================================================================== 2. The answer
s = new_slide(title="answer")
header(s, "THE ANSWER", "One way of deciding, used three times: the plan, the Board, the HQ")
add_text(s, 0.67, 1.4, 11.8, 0.65, [[("The same mechanism runs all three. ", True), ("The rule is published before the argument. Options go "
         "to one named decision owner, on a date. The decision, the rationale and the dissent are recorded. Only exceptions come back up.", False)]],
         size=12.5, color=INK, where="through-line")
tiles = [("DELIVERABLE 1  //  PLANNING", "$25M in two releases",
          "Every ask needs a business case. $15M on 1 January, $10M at Q1 close if the numbers hold. No flat cuts."),
         ("DELIVERABLE 2  //  BOARD", "One test, one decision owner",
          "Both executives' positions go into the memo; the CEO decides at the lock; the dissent is recorded. The book ships 7 days before the Board."),
         ("DELIVERABLE 3  //  HQ", "Four decision gates in 120 days",
          "Can we, what may we offer, who may bind Nu, do we commit: each the CFO's call on a pack I assemble. Licensed executives sign."),
         ("DELIVERABLE 4  //  LEADERSHIP", "Rigor inside the business",
          "I work in the numbers with the GMs. Experts own their judgment, I own the process, the CFO owns the decision and its risk.")]
for i, (lab, head, body) in enumerate(tiles):
    card(s, 0.67 + i * 3.0, 2.2, 2.8, 2.4, lab, head, [body], body_size=12.5)
band(s, 0.67, 4.8, 11.8, 1.5, [[("THE ASK  ", True, LILAC_TXT)],
     [("Approve the $25M FY27 cap in two releases, the Board calendar and the HQ charter. Up to $15M more is held in "
       "reserve, outside the plan: the CFO decides, and Risk can block it. I report back in 90 days.", True)]], size=15)
notes(s, "One way of deciding, used three times. Three decisions, each with a number. The reserve is stated up front, with its conditions.")

# ================================================================== 3. Assumptions (short form)
s = new_slide(title="assumptions")
header(s, "WHAT I KNOW AND WHAT I AM ASSUMING", "What I know, and the assumptions that could break the plan")
rows = [["#", "Statement", "Class", "Basis"],
        ["A1", "Nu Q2'26: 138.9M customers; monthly ARPAC $17.1; cost to serve $1.0; efficiency ratio 19.5%; ROE 33%; net income $1.06B.", "FACT", "Q2'26 release and 6-K"],
        ["A2", "Management expects the FY26 efficiency ratio near 20% and will keep the US entry within about 100 bps of efficiency-ratio impact.", "FACT (management guidance)", "Q2'26 earnings call, 13 Aug 2026"],
        ["A4", "The two MTeam members disagree about the growth-opex envelope, core markets against new bets, not about regulatory capital by entity. If wrong, slide 7's capital ranking is the decision.", "ASSUMPTION", "Case brief is silent"],
        ["A5", "$8M of the asks can be met on the shared platform at near-zero marginal cost; what Engineering cannot confirm moves to the case track, never above the cap.", "ASSUMPTION", "My estimate"],
        ["A6", "Spend above $25M only from H1 profit above plan, inside the ~20% guide, with NPL 90+ on plan. The CFO decides; Risk can block; IR sets the disclosure.", "RULE (proposed)", "To agree with the CFO"],
        ["A10", "The cap is FY27's: first release on 1 Jan 2027, second at Q1 close (31 Mar), the reserve at H1 close. The hub needs 4-6 people.", "ASSUMPTION", "Calendar fiscal year; sizing is mine"],
        ["A12", "Net new FY27 cost for the hub and the HQ is small and is funded inside the $25M before the GM releases. If it is material, the second release shrinks by that amount.", "ASSUMPTION", "To size with the CFO this week"]]
bottom = table(s, 0.67, 1.7, [0.55, 7.35, 1.95, 2.15], rows, size=10.5)
footer(s, "Confirming now: the real MTeam positions; the Board and committee calendar; who owns the planning model; local-CFO "
          "decision rights in Mexico and Colombia; the HQ project's status and whether a site shortlist exists; today's cycle time and "
          "forecast accuracy. Full list A1-A13: backup A.", y=max(bottom + 0.15, 6.35), h=0.6)
notes(s, "A4 matters most: if the dispute is about capital by entity, slide 7 carries the decision.")

# ================================================================== 4. D1 assessment, options and principle
s = new_slide(title="d1-principle")
header(s, "DELIVERABLE 1  //  PLANNING", "Staged release beats flat cuts: the precedent matters more than the $35M")
rect(s, 0.67, 1.75, 3.4, 4.8, fill=DARK, line=None, adj=0.05)
add_text(s, 0.8, 1.9, 3.15, 0.3, ["THE GAP"], size=10, bold=True, color=LILAC_TXT, align=PP_ALIGN.CENTER, where="gap label")
add_text(s, 0.8, 2.2, 3.15, 1.05, ["$35M"], size=58, bold=True, color=WHITE, align=PP_ALIGN.CENTER, where="gap")
add_text(s, 0.85, 3.3, 3.05, 1.55,
         ["$60M of asks against a $25M cap.", "About 15 bps of annualized revenue.", "About 1.1% of annual opex.",
          "About 0.2 points of efficiency ratio."], size=12.5, color=SOFT, align=PP_ALIGN.CENTER, space_after=[6, 3, 3, 3],
         where="gap facts")
add_text(s, 0.85, 4.95, 3.05, 1.5,
         ["Efficiency ratio 19.5% in Q2'26 (17.6% in Q1); FY26 expected near 20%. Cost to serve $1.0 a month "
          "($0.8 a year earlier) against $17.1 of ARPAC."], size=10.5, color=SOFT, align=PP_ALIGN.CENTER, where="gap context")
add_text(s, 4.3, 1.7, 8.2, 1.2, [[("What is broken today", True)],
         "Asks arrive as one $60M number, with no shared template, no shared data and a cap set top-down. So the argument is "
         "about the number, never the evidence. This first central cycle decides whether GMs bring real cases next year or pad their asks."],
         sizes=[14, 12], colors=[PURPLE, BODY], space_after=[3, 0], where="broken today")
rows = [["Three ways to close the gap", "Score", "Why"],
        [[("Equal cut", True), (" (every unit gets 42% of its ask)", False)], "57.8", "Fast, and it teaches GMs to pad next year's ask"],
        [[("Core markets first", True), (" (pause new bets)", False)], "63.8", "Honest, but new bets get nothing to earn"],
        [[("Business-case readiness and staged release", True), (" (recommended)", False)], [("91.8", True)],
         "Wins on growth return, ownership and agility; loses only on speed"]]
bottom = table(s, 4.3, 3.0, [3.7, 0.8, 3.7], rows, size=11, header_size=11, row_min=0.42, highlight_last=True, max_bottom=5.0)
add_text(s, 4.3, bottom + 0.15, 8.2, 1.45, [[("Budgets should work like Nu's credit", True)],
         "Nu's 20-F describes low-and-grow: small limits, raised on behavior. Budgets get the same test-and-learn: start small, "
         "grow when four indicators hold: activation, 30-day retention, CAC and unit cost against plan. That is smart efficiency, run by owners."],
         sizes=[14, 12], colors=[PURPLE, BODY], space_after=[3, 0], where="principle")
footer(s, "Scores and weights are my judgment (A9); option 3 stays first if any single weight moves 10 points (backup C). Sources: Nu "
          "Q2'26 release and 6-K (opex $806.2M; net revenue ~$4.1B a quarter), Q2'26 call, Form 20-F. Calculations: the 15 bps is on "
          "managerial revenue, ~$23.5B annualized; the 1.1% is on opex, ~$3.2B annualized; the efficiency-ratio points are on Nu's "
          "denominator (NII plus fees).", y=6.7, h=0.65)
notes(s, "What is broken is the process, not the number. $35M is about 0.2 points of efficiency ratio; $25M is about 0.15.")

# ================================================================== 5. Bridge, releases and reserve
s = new_slide(title="bridge")
header(s, "DELIVERABLE 1  //  PLANNING", "Five answers for $60M: fund now, stage, reallocate, defer or stop")
cd = CategoryChartData()
cd.categories = ["GM asks", "Reallocate: shared platform", "Stage: 60-day case track", "Defer: H2 reserve, not in plan", "Fund now and stage: two releases"]
cd.add_series("base", (0, 52, 40, 25, 0))
cd.add_series("value", (60, 8, 12, 15, 25))
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED, Inches(0.5), Inches(1.75), Inches(6.6), Inches(4.2), cd)
ch = gf.chart
ch.has_legend = False
ch.has_title = False
ca = ch.category_axis
ca.reverse_order = True
ca.tick_labels.font.size = Pt(11)
ca.tick_labels.font.name = FONT
ca.tick_labels.font.color.rgb = INK
ca.format.line.fill.background()
va = ch.value_axis
va.visible = False
va.has_major_gridlines = False
va.maximum_scale = 61
plot = ch.plots[0]
plot.gap_width = 55
plot.overlap = 100
base_s, val_s = plot.series
base_s.format.fill.background()
base_s.format.line.fill.background()
colors = [RGBColor(0x5B, 0x4B, 0x7A), RGBColor(0x8B, 0x5C, 0xF6), RGBColor(0x8B, 0x5C, 0xF6), ORANGE, PURPLE]
labels = ["$60M", "$8M", "$12M", "$15M", "$25M"]
for i, pt in enumerate(val_s.points):
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = colors[i]
    dl = pt.data_label
    dl.has_text_frame = True
    dl.text_frame.text = labels[i]
    dl.position = XL_LABEL_POSITION.INSIDE_END
    r = dl.text_frame.paragraphs[0].runs[0]
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.name = FONT
    r.font.color.rgb = WHITE
right = [("FUND NOW  //  FIRST RELEASE, 1 JAN 2027", "$15M", "$10.5M to core markets (70%) and $4.5M to new bets (30%), for cases already vetted."),
         ("STAGE-GATE  //  SECOND RELEASE, 31 MAR", "$10M", "To units whose Q1 indicators hold, decided on Q1 data in early April. Completes the $25M."),
         ("DEFER  //  H2 RESERVE, NOT IN PLAN", "Up to $15M", "Only from H1 profit above plan, inside the ~20% guide, with NPL 90+ on plan. "
          "The CFO decides; Risk can block; IR sets the disclosure.")]
yy = 1.8
for lab, head, body in right:
    h = 1.55
    card(s, 7.3, yy, 5.2, h, lab, head, [body], label_color=(ORANGE if "RESERVE" in lab else PURPLE), body_size=12,
         heading_size=16)
    yy += h + 0.1
footer(s, "$8M of country asks reuse the shared code base at near-zero marginal cost (A5); anything Engineering cannot confirm "
          "moves to the case track, and the cap does not grow. $12M gets 60 days to a case with my team, then is funded, deferred or "
          "stopped. Stop is real: a unit below plan on two of four indicators has its second-release share held (slide 15).",
       y=6.65, h=0.65)
notes(s, "Every dollar gets one of five answers: fund now, stage, reallocate, defer or stop; the reserve opens only on its conditions.")

# ================================================================== 6. Hub, capacity, cadence and D1 risks
s = new_slide(title="hub")
header(s, "DELIVERABLE 1  //  PLANNING", "A lean hub and three meetings bring planning in-house")
card(s, 0.67, 1.75, 4.6, 3.45, "THE HUB", "A planning hub of 4-6 people (A10)",
     ["One planning model, read straight from Nu's data platform.",
      "Its KPIs: cycle time under 40 days, forecast accuracy, and 100% of spend with an approved case.",
      "Finance partners in each market and vertical: solid line to the hub, dotted line to the GM, in place by late November.",
      "A PMO agent and a review agent do the consolidation (backup H). I hire for modeling and plain-English writing."],
     body_size=11.5)
rows = [["Meeting", "Who runs it", "Output"],
        [[("Squad clinics", True), (", weekly", False)], "Finance partner, with the GM", "A case in the standard template, co-signed by the GM"],
        [[("GM one-to-ones", True), (", every two weeks", False)], "Me", "No surprises at the lock; every ask has a path to yes"],
        [[("Release review", True), (", quarterly", False)], "CFO decides, on my recommendation", "Release memo to the GMs the same day"]]
bottom = table(s, 5.5, 1.75, [2.15, 2.25, 2.6], rows, size=10.5, header_size=11, row_min=0.42, max_bottom=3.75)
band(s, 5.5, bottom + 0.12, 7.0, 5.2 - bottom - 0.12,
     [[("Capacity, not only money: ", True), ("every case states the headcount and engineering capacity it needs, and Engineering "
       "confirms it at the lock, so a release never funds work the platform cannot staff. Partners and GMs own the case and its "
       "numbers; I assemble the pack and recommend; the CFO decides; the hub runs the release.", False)]], size=11.5, dark=False)
risks6 = [("RISK", "GMs game the indicators", "Finance partners own the read-out; the review agent ties every figure to the data."),
          ("RISK", "The $8M platform estimate slips", "What Engineering cannot confirm moves to the case track; the cap does not grow."),
          ("RISK", "The hub becomes a cost GMs resent", "Net new cost is funded inside the cap and shown first (A12). If cycle time slips, simplify; do not hire.")]
for i, (lab, head, body) in enumerate(risks6):
    card(s, 0.67 + i * 3.95, 5.38, 3.85, 1.3, lab, head, [body], label_color=ORANGE, body_size=10.5, heading_size=12.5)
notes(s, "Four to six people, one model, three meetings; every ask states its people and engineering capacity; three ways it fails, each with an answer.")

# ================================================================== 7. D2 capital
s = new_slide(title="capital")
header(s, "DELIVERABLE 2  //  C-SUITE AND BOARD", "The capital dispute becomes one choice for the CEO: four uses, one test")
cols = [("WHAT IS ON THE TABLE", "Four uses of equity",
         ["Brazil deepening.",
          "Nu Mexico: bank authorized 10 Jul 2026, operating since 6 Aug; $4.2B planned to 2030.",
          "US charter: OCC conditional approval 29 Jan 2026; capital in by about 29 Jan 2027.",
          "Buyback of up to $1.0B, announced Jun 2026."]),
        ("THE TEST", "Return on allocated capital",
         ["After-tax return on the extra equity each use needs, against that market's cost of equity, with headroom by entity "
          "in a downside case.",
          "Beating the cost of equity creates value even if it dilutes the 33% group ROE. The buyback is the bar."]),
        ("WHO DECIDES, BY WHEN", "Before Investor Day",
         ["The CFO recommends, the CEO decides, the Board approves at this meeting, the last before Investor Day on 8 Dec.",
          "One page per use: return, capital, headroom, risk. Cost of equity comes from Treasury."])]
for i, (lab, head, body) in enumerate(cols):
    card(s, 0.67 + i * 4.0, 1.75, 3.8, 3.35, lab, head, body, body_size=11.5)
band(s, 0.67, 5.3, 11.8, 1.3,
     [[("On opex, both positions go into the T-16 memo as options: ", True), ("core markets first, new bets first, or the staged "
       "70/30 release. I recommend 70/30: mature cohorts earn about $31 a month at eight years against $17.1 for the group, and new bets "
       "are judged on activation, retention and unit cost, not first-quarter profit. The CEO decides; the dissent is recorded; the Board "
       "sees one recommendation. If the dispute is about equity, the test above decides it (A4).", False)]],
     size=11.5)
footer(s, "Sources: Nu Q2'26 release (ROE 33%, ARPAC $17.1); Q1'26 cohort slide ($30.8 at 96 months vs $15.9 total); Nu "
          "announcements (Mexico 10 Jul and 6 Aug 2026; OCC 29 Jan 2026; buyback 4 Jun 2026); Nu IR (Investor Day 8 Dec 2026).",
       y=6.75, h=0.45)
notes(s, "Equity by entity first. On opex, both positions go in the memo as options; I recommend 70/30; the CEO decides and the dissent is recorded.")

# ================================================================== 8. Pushback
s = new_slide(title="pushback")
header(s, "DELIVERABLE 2  //  C-SUITE AND BOARD", "When a senior GM pushes back: four moves, in the same order every time")
steps = [("1", "Talk privately first", "I confirm the strategic goal before any budget line comes up. Nobody is corrected in front of the CFO."),
         ("2", "Move to the evidence", "From \"I need $25M\" to \"what would prove this works?\" The published rule does the arguing, not me."),
         ("3", "Offer a conditional yes", "A funded start, the $4.5M for new bets in the first release, and a reserved claim on the second release once the indicators hold."),
         ("4", "Present together", "The GM and I present the plan together at the lock. The CEO's office is briefed in advance and sees the same rule.")]
for i, (n, head, body) in enumerate(steps):
    card(s, 0.67 + i * 3.0, 1.8, 2.8, 2.95, f"STEP {n}", head, [body], body_size=12.5)
band(s, 0.67, 4.95, 11.8, 1.45, [[("Why it works without line authority: ", True), ("people accept a decision they argued "
     "against when the criteria were public before it was made. The rule is public on day one, the conversation is private, "
     "and the sign-off is joint.", False)],
     [("With the two MTeam members: ", True), ("the same four moves run separately at T-21, before they meet at the lock.", False)]],
     size=12.5)
footer(s, "Fair process: Kim and Mauborgne, \"Fair Process: Managing in the Knowledge Economy\", Harvard Business Review.", y=6.6, h=0.35)
notes(s, "The same four moves every time; the GM keeps the goal and earns the next release.")

# ================================================================== 9. Board rhythm, diagnosis and D2 risks
s = new_slide(title="board")
header(s, "DELIVERABLE 2  //  C-SUITE AND BOARD", "Every quarter: the book goes out 7 days before the Board, after a lock at 16 days")
add_text(s, 0.67, 1.6, 11.8, 0.62, [[("Why the book is late: ", True), ("no locked outline, no standard data tables, and units building "
         "rival decks. The fix starts today, with the Corporate Secretary, and the same countdown then runs for every Board from the "
         "Secretary's calendar. That calendar is the one the ask approves.", False)]], size=11.5, color=INK, where="diagnosis")
weeks = [("T-28 TO T-21", "Outline locked", "Today I issue the locked outline and standard data tables, so units stop building rival decks."),
         ("T-20 TO T-14", "Executive lock", "One-to-ones at T-21. The capital options memo settles the ranking and the 70/30 split with the CEO and CFO at T-16."),
         ("T-13 TO T-7", "One pen", "Inputs freeze at T-10. Committee chairs are briefed at T-9. The book is on the portal at T-7."),
         ("T-6 TO T-0", "Decisions, not read-outs", "Directors' questions are answered in writing by T-3, so the meeting goes to the capital plan.")]
for i, (lab, head, body) in enumerate(weeks):
    card(s, 0.67 + i * 3.0, 2.27, 2.8, 1.85, lab, head, [body], body_size=11)
band(s, 0.67, 4.24, 11.8, 0.8, [[("Two items cannot slip: ", True), ("the capital plan, because this is the last Board before "
     "Investor Day on 8 Dec, and the statutory items for new entities: open an entity only when a license or tax substance "
     "requires it; officers and signing follow the stage (backup B).", False)]], size=11.5, dark=False)
risks9 = [("RISK", "Still no agreement at the lock", "Both positions go in the memo; the CEO decides; the dissent is recorded, so the Board sees one recommendation."),
          ("RISK", "Materials are still late at T-10", "The locked outline decides: anything outside it goes to the appendix or the next meeting."),
          ("RISK", "The book holds non-public Q3 results", "Portal only; IR sets what directors may discuss. The review agent's drafts run on Nu-approved tools, and a named owner signs each answer.")]
for i, (lab, head, body) in enumerate(risks9):
    card(s, 0.67 + i * 3.95, 5.15, 3.85, 1.5, lab, head, [body], label_color=ORANGE, body_size=10.5, heading_size=12)
footer(s, "Board assumed on Tue 17 Nov, four weeks from today (A7). Q3'26 results are due 12 Nov.", y=6.8, h=0.35)
notes(s, "The book is late because nothing is locked. If the two executives still disagree at the lock, the CEO decides and the dissent is recorded.")

# ================================================================== 10. D3 assessment and site choice
s = new_slide(title="hq-assess")
header(s, "DELIVERABLE 3  //  HQ PROJECT", "The HQ is a governance project before it is a real-estate deal")
card(s, 0.67, 1.75, 5.7, 2.55, "THE ASSESSMENT", "What can go wrong before a lease exists",
     ["A regional hub and campus, negotiated across Legal, Tax, Compliance, People and Workplace, in more than one "
      "jurisdiction, with no line over any of them. The jurisdiction is not given.",
      "Three early failures: tax residence moves or a taxable presence appears; the incentive is worth less once Pillar Two "
      "applies; an officer signs before being approved and insured. So: governance first, campus last."],
     body_size=11)
card(s, 0.67, 4.45, 5.7, 2.15, "WHAT IS NEGOTIATED", "Two kinds of agreement",
     ["Local regulatory agreements: the incentive or grant, any tax ruling, registrations and permits, work permits "
      "(Legal, Tax, Compliance, People).",
      "Campus infrastructure agreements: lease or build-to-suit, fit-out, utilities, land and zoning (Workplace, Legal)."],
     body_size=11)
card(s, 6.55, 1.75, 5.95, 4.85, "HOW I WOULD CHOOSE THE SITE", "Four hard tests, then five criteria",
     ["1. Regulatory: the hub's activities need no license Nu does not have (Compliance).",
      "2. Tax: people and decisions there do not move Nu Holdings' tax residence or create a taxable presence (Tax opinion at gate 1).",
      "3. Officers: Nu can field directors who pass fit-and-proper (Legal, People).",
      "4. Substance: real payroll and assets, so the incentive survives Pillar Two (Tax, People).",
      "Then five weighted criteria: talent and wage cost; incentive net of clawback; campus cost and timing; modeled tax rate; "
      "political and reputational risk (weights in backup J).",
      "I start where Nu already has licensed substance and a local CFO (Brazil, Mexico, Colombia) or is building (the US). "
      "Two sites go to term sheet; the CEO and CFO pick. If a shortlist already exists (A13), the tests run on it at gate 1 "
      "and the criteria become the tie-break."],
     body_size=11)
footer(s, "Nu Holdings is Cayman-incorporated (Form 20-F). Nu is above the EUR 750M Pillar Two threshold; 15% minimum; 2026 carve-out "
          "9.4% of payroll and 7.4% of tangible assets (OECD GloBE, Art. 5.3, 9.2). Brazil applies a domestic minimum top-up tax "
          "from 2025. Jurisdiction open (A13); no savings figure until a site is chosen.", y=6.72, h=0.6)
notes(s, "A governance project before a real-estate deal: hard tests pick what can be considered, criteria rank it, the CEO and CFO pick.")

# ================================================================== 11. D3 gates with negotiation and walk-aways
s = new_slide(title="hq-gates")
header(s, "DELIVERABLE 3  //  HQ PROJECT", "Four decision gates in 120 days; nothing is signed between them")
rows = [["Gate", "Owner", "What we negotiate", "Exit test", "Walk-away"],
        [[("Gate 1, D+30: can we?", True), ("\nTax and residence", False)], "Tax, with Compliance", "Nothing yet: internal",
         "Tax-residence and taxable-presence opinion for both sites; modeled tax rate holds under Pillar Two; Compliance confirms permission",
         "Any site where residence or a taxable presence moves, whatever the incentive"],
        [[("Gate 2, D+60: what may we offer?", True), ("\nRegulatory and grant terms", False)], "Legal, with Tax",
         "The incentive agreement and any ruling; commitments only inside the published plan",
         "Signed term sheet: 85% of headcount keeps 85% of the grant; 24 months' grace; one site dropped",
         "A cliff-edge clawback; a headcount commitment above plan; terms that bind other markets"],
        [[("Gate 3, D+90: who may bind Nu?", True), ("\nOfficers and signing", False)], "Legal and People", "Nothing external: Nu's own governance",
         "Officers named from licensed executives; fit-and-proper files in (started at gate 1); D&O bound; signing matrix and powers of attorney approved",
         "Any appointment without D&O cover or regulator approval"],
        [[("Gate 4, D+120: do we commit?", True), ("\nLease and campus", False)], "Workplace, with Legal",
         "Lease or build-to-suit, fit-out, utilities, infrastructure contribution",
         "Lease inside the FY27 plan, cost phased, a break option matched to the grant's grace, a long-stop date tied to gate 2",
         "No break option, or cost committed before the incentive is signed"]]
bottom = table(s, 0.67, 1.7, [1.75, 1.45, 2.6, 3.4, 2.6], rows, size=10.5, header_size=11, row_min=0.5, max_bottom=5.75)
band(s, 0.67, bottom + 0.15, 11.8, 1.2,
     [[("Gate 0 is the charter signed today: ", True), ("why, the outcomes, the boundaries, the CFO as sponsor. Each gate after it "
       "is a decision the CFO makes on a pack I assemble; the walk-aways are the red lines of the negotiating mandate. Legal leads the "
       "room. What we give stays inside the published plan: no headcount or spend commitment the $25M cap cannot carry.", False)]],
     size=11.5)
footer(s, "Dates run from the charter signed today (D+0 = 20 Oct): gate 1 on 19 Nov, two days after the Board; gates 2-4 in mid-December, "
          "mid-January and mid-February (backup F). The 85% and 24-month terms are proposals. Q2'26 showed how real-estate timing "
          "moves the efficiency ratio, so the lease cost is phased.", y=6.72, h=0.6)
notes(s, "Gate 0 is the charter. Four decision gates: can we, what may we offer, who may bind Nu, do we commit. The CFO decides each on my pack; nothing is signed between them.")

# ================================================================== 12. D3 statutory officers and interim leadership
s = new_slide(title="hq-officers")
header(s, "DELIVERABLE 3  //  HQ PROJECT", "Experts own the judgment, I own the process, the CFO owns residual risk")
card(s, 0.67, 1.75, 5.8, 3.6, "STATUTORY OFFICER IMPLICATIONS", "Who signs, who sits, who carries the risk",
     ["Signing: central only until gate 3, which also avoids an accidental taxable presence; then dual signature above limits "
      "set by category, and powers of attorney limited in scope and time, reviewed annually.",
      "Seats: licensed executives who pass the regulator's fit-and-proper test. Not me: I run the process and hold no power of attorney.",
      "Liability: officers carry personal liability for the entity's filings and conduct (Legal confirms per site). D&O names "
      "the new entity before activation. Regulator approval sets the critical path, so files start at gate 1.",
      "Capital: a regulated hub entity needs its own capital plan and goes on the Board's statutory agenda (backup B)."],
     body_size=10.5)
card(s, 6.67, 1.75, 5.8, 3.6, "LEADING AS AN INTERIM ACROSS FIVE FUNCTIONS", "One pod, one log, clear decision rights",
     ["Weekly 45-minute pod: Legal, Tax, Compliance, People and Workplace, each with one lead who can commit their function.",
      "I own the integrated outcome and the decision process; I do not second-guess a tax opinion. Each function owns its "
      "technical judgment. The CFO owns acceptance of any material residual risk at each gate; the CEO decides the site; the "
      "entity's board appoints officers; the Audit and Risk Committee hears at gates 1 and 3.",
      "One decision log, kept by the PMO agent; escalation to the CFO within 48 hours; a named deputy covers me.",
      "When functions disagree: their words go in the log; 48 hours to a joint view; if none, both options go to the CFO with "
      "mine. Legal and Compliance can veto on legality; nobody vetoes on price."],
     body_size=10.5)
risks12 = [("RISK", "Tax residence moves", "Opinion at gate 1; meeting patterns unchanged; no early signing (Tax)."),
           ("RISK", "Officer approval runs late", "Files at gate 1; no activation before approval; the lease long-stop moves with it (Legal, People)."),
           ("RISK", "Dependence on me", "A named deputy, the decision log, and gate packs the CFO can decide on without me.")]
for i, (lab, head, body) in enumerate(risks12):
    card(s, 0.67 + i * 3.95, 5.47, 3.85, 1.25, lab, head, [body], label_color=ORANGE, body_size=10.5, heading_size=12)
footer(s, "Signing limits and the delegation model are proposals. If the hub sits inside an existing licensed entity, gate 3 shrinks to a "
          "powers-of-attorney update (backup B). Full HQ risk register and site scorecard: backup J.", y=6.85, h=0.45)
notes(s, "Experts own the judgment, I own the process, the CFO owns the residual risk. Licensed executives sign; I hold no power of attorney. Disagreements get 48 hours, then both options go to the CFO with mine.")

# ================================================================== 13. D4 three rules
s = new_slide(title="leadership")
header(s, "DELIVERABLE 4  //  LEADERSHIP STYLE AND OPERATING PHILOSOPHY", "Influence with hands-on rigor: a few clear rules, published early")
rules = [("Publish the rule, then negotiate",
          ["People accept outcomes they argued against when the criteria were public first.",
           "I publish the scoring on day one and do not change it mid-cycle. Every GM sees the same four indicators."]),
         ("Fix the numbers in private",
          ["Model gaps are closed one to one before anything reaches the CFO or the Board.",
           "Senior people trust you when you never surprise them in a room."]),
         ("Own the model and the clauses",
          ["I build the model and read the clauses myself. AI does the first pass; I still do the last.",
           "Influence at this level is earned by being right on the detail."])]
for i, (head, body) in enumerate(rules):
    card(s, 0.67 + i * 4.0, 1.8, 3.8, 2.75, f"RULE {i + 1}", head, body, body_size=12.5)
band(s, 0.67, 4.8, 11.8, 1.55, [[("Where I push back on the CFO: ", True), ("when a decision skips the published rule. ", False)],
     [("What I leave alone for 90 days: ", True), ("local-CFO decision rights in Mexico and Colombia, credit-model governance, "
       "and the IR calendar.", False)]], size=13)
footer(s, "Fair process: Kim and Mauborgne (HBR).", y=6.55, h=0.35)
notes(s, "Three rules, each tied to a mechanism shown earlier.")

# ================================================================== 14. D4 balance, problem by problem
s = new_slide(title="balance")
header(s, "DELIVERABLE 4  //  LEADERSHIP STYLE AND OPERATING PHILOSOPHY", "How I work in a matrix: what I own, what experts own, what I leave alone")
rows = [["Problem", "I own: the model, the memo, the log", "I influence, with no line authority", "Others own: the decision, the signature"],
        ["Planning", "Build and review the single model; write the case template; read every first-release case before the lock",
         "GM one-to-ones; the lock with the CEO and CFO; the release memo", "Deciding releases (the CFO does); local-CFO decision rights"],
        ["Board", "Draft the capital options memo; tie every figure in the book to its source; answer directors' data questions by T-3",
         "The two executives at T-21; committee chairs at T-9", "Presenting the ranking (the CFO does); the IR calendar"],
        ["HQ", "Read the term sheets and clauses; keep the decision log; build the site scorecard",
         "The weekly pod; the mandate with the CFO; the CEO on the site", "Signing; holding a power of attorney; sitting as a statutory officer"]]
bottom = table(s, 0.67, 1.75, [1.5, 4.0, 3.15, 3.15], rows, size=11.5, header_size=11.5, row_min=0.75, max_bottom=5.2)
band(s, 0.67, bottom + 0.25, 11.8, 1.45,
     [[("The same split in all three problems. ", True), ("I own the process and the integrated recommendation. Experts own their "
       "technical judgment. The accountable executive owns the decision and its residual risk.", False)],
      [("The team I build: ", True), ("I hire for modeling and plain-English writing; each hub member owns one market's release pack; "
       "a named deputy sits in the HQ pod, so by the second release nothing depends on me alone.", False)]], size=12)
notes(s, "The same split in all three problems: I own the model, the memo and the log; experts own their judgment; the accountable executive owns the decision and the signature.")

# ================================================================== 15. Next 90 days, stop rules and the ask
s = new_slide(title="ask")
header(s, "NEXT 90 DAYS AND THE ASK", "Next 90 days: the Board on 17 Nov, the first release on 1 Jan, and the ask")
phases = [("NOW TO THE BOARD", "20 Oct to 19 Nov",
           ["Today: Board outline out, HQ charter signed, tax gate launched.",
            "Early Nov (T-16): capital options lock with the CEO, CFO and MTeam.",
            "17 Nov: Board approves the FY27 plan. 19 Nov: HQ gate 1."]),
          ("BOARD TO YEAR-END", "Build the hub",
           ["Late Nov: hub in place, partners embedded; PMO agent live, status templates retired.",
            "Early Dec: squad clinics running; $12M case track under way.",
            "Mid-Dec (D+60): HQ grant terms agreed."]),
          ("JANUARY TO MARCH", "Open FY27",
           ["1 Jan: FY27 starts; $15M first release.",
            "Mid-Jan (D+90): HQ officers in place; my 90-day report to the CFO.",
            "Mid-Feb: lease (D+120). 31 Mar: Q1 close; second release decided in early April."])]
for i, (lab, head, body) in enumerate(phases):
    card(s, 0.67 + i * 4.0, 1.75, 3.8, 2.3, lab, head, body, body_size=11)
band(s, 0.67, 4.2, 11.8, 0.95, [[("Stop rules, public from today: ", True), ("a unit below plan on two of four indicators has its "
     "second-release share held. Efficiency ratio above the guide in force at Q1 close: the second release is held for every unit. "
     "Credit cost or NPL 90+ above plan: the reserve stays shut, and Risk can block it.", False)]], size=11.5, dark=False)
band(s, 0.67, 5.3, 11.8, 1.35, [[("THE ASK  ", True, LILAC_TXT)],
     [("Approve the $25M FY27 cap in two releases, the Board calendar and the HQ charter.", True)],
     [("In 90 days (mid-January): every first-release dollar has an approved case; the Board book went out 7 days early; HQ gates 1 "
       "and 2 are cleared. You are approving one way of deciding, used three times, and I report on it in 90 days.", False)]], size=12.5)
notes(s, "Three dates, three stop rules, the ask: one way of deciding, used three times, reported on in 90 days.")

# ================================================================== BACKUP divider
s = new_slide(title="backup")
add_text(s, 0.67, 2.6, 12.0, 1.0, ["Backup"], size=40, bold=True, color=INK, align=PP_ALIGN.CENTER, where="backup")
add_text(s, 1.2, 3.7, 10.9, 1.6, ["A  Full assumptions (A1-A13)   |   B  Statutory entity model   |   C  Option sensitivity",
                                  "D  Sources   |   E  Full-year split by horizon   |   F  Day-by-day timeline   |   G  Stakeholder map",
                                  "H  Two AI agents   |   I  Scenarios and stop rules   |   J  HQ site scorecard and risk register"],
         size=14, color=BODY, align=PP_ALIGN.CENTER, where="backup list")
notes(s, "Not presented.")

# ================================================================== A. Full assumptions
s = new_slide(title="A")
header(s, "BACKUP A", "All assumptions, with the facts they rest on")
rows = [["#", "Statement", "Class", "Basis"],
        ["A1", "Nu Q2'26: 138.9M customers, monthly ARPAC $17.1, cost to serve $1.0, efficiency ratio 19.5%, ROE 33%, net income $1.06B.", "FACT", "Q2'26 release and 6-K"],
        ["A2", "FY26 efficiency ratio near 20%; the US entry kept within about 100 bps of efficiency-ratio impact.", "FACT (management guidance)", "Q2'26 call, 13 Aug 2026"],
        ["A3", "The $60M asks and $25M cap are the case's figures. $35M is about 15 bps of annualized revenue, about 1.1% of opex and about 0.2 points of efficiency ratio ($25M: about 0.15).", "ASSUMPTION + CALCULATION", "Case brief; Q2'26 figures"],
        ["A4", "The MTeam dispute is about the growth-opex envelope (core vs new bets), not regulatory capital by entity.", "ASSUMPTION", "Case brief is silent"],
        ["A5", "$8M met on the shared platform at near-zero marginal cost; unconfirmed scope moves to the case track; $12M needs a case, ready in 60 days.", "ASSUMPTION", "My estimate; Engineering confirms"],
        ["A6", "Spend above $25M only from H1 profit above plan, inside the ~20% guide, NPL 90+ on plan; CFO decides, Risk can block, IR sets disclosure.", "RULE (proposed)", "To agree with the CFO"],
        ["A7", "Board in 4 weeks (assumed Tue 17 Nov); Investor Day 8 Dec 2026; the Board approves the capital plan before the Investor Day story is set.", "FACT (dates) + ASSUMPTION", "Case brief; Nu IR"],
        ["A8", "HQ jurisdiction open. Nu is above the EUR 750M Pillar Two threshold; 2026 carve-out 9.4% of payroll, 7.4% of tangible assets.", "FACT (rule) + ASSUMPTION (site)", "OECD GloBE rules"],
        ["A9", "Option scores and weights are my judgment; third-party benchmarks are context only.", "ASSUMPTION", "Slide 4; backups C, D"],
        ["A10", "FY27 envelope: first release 1 Jan 2027, second at Q1 close (31 Mar), reserve at H1 close. Hub of 4-6 people.", "ASSUMPTION", "Calendar fiscal year; sizing is mine"],
        ["A11", "The hub can run Nu-approved AI tools inside Nu's environment; Security and Risk approve data access before the agents go live.", "ASSUMPTION", "To confirm with Security and Risk"],
        ["A12", "Net new FY27 cost for the hub and the HQ is small and funded inside the $25M before the GM releases; if material, the second release shrinks by that amount.", "ASSUMPTION", "To size with the CFO"],
        ["A13", "HQ jurisdiction open; two candidate sites go to term sheet; the site scorecard weights are my judgment (backup J).", "ASSUMPTION", "Case brief is silent"]]
table(s, 0.67, 1.4, [0.55, 7.55, 1.9, 2.0], rows, size=9.5, header_size=10, row_min=0.3, max_bottom=7.05)
notes(s, "Backup.")

# ================================================================== B. Statutory entity model
s = new_slide(title="B")
header(s, "BACKUP B", "Open an entity only when a license or tax substance requires it; officers and signing follow the stage")
rows = [["Stage", "Trigger for a local entity", "Setup", "Officers, signing and insurance"],
        [[("1. Market test", True), (" (new countries)", False)], "No local license needed and no taxable presence; served from an existing licensed entity or a partner bank (as the US is today, through Lead Bank, since 10 Sep 2026).", "No local entity. Cross-border or partner model.", "Central signing only. No local director. Prevents an accidental permanent establishment."],
        [[("2. Licensed scale-up", True), (" (US charter, new verticals)", False)], "A regulator requires a local licensed entity (deposits, lending, custody), or substance rules require local payroll and assets.", "Dedicated licensed entity with its own capital plan and headroom.", "Fit-and-proper directors; two central officers plus one local leader; dual signature above a set limit; D&O before activation."],
        [[("3. Core hub", True), (" (Mexico, a bank since Aug 2026; Colombia; the HQ)", False)], "Full banking or fintech license in force; local CFO in place (already so in Mexico and Colombia).", "Full licensed entity; regional campus where the HQ lands.", "Audit and Risk Committee oversight; dual signature above limits set by category; director liability insurance; annual review of powers of attorney. If the HQ sits inside an existing licensed entity: a powers-of-attorney update only."]]
bottom = table(s, 0.67, 1.75, [2.0, 3.8, 2.4, 3.8], rows, size=11, row_min=0.5)
footer(s, "Sources: Nu announcements (US launch via Lead Bank, 10 Sep 2026; OCC conditional approval, 29 Jan 2026; Mexico bank "
          "authorization, 10 Jul 2026; local finance leadership per the CFO appointment release). Signing limits are proposals.",
       y=max(bottom + 0.2, 6.3), h=0.5)
notes(s, "Backup.")

# ================================================================== C. Sensitivity
s = new_slide(title="C")
header(s, "BACKUP C", "Option 3 stays first under every 10-point change to a single weight")
W = {"cap": .30, "ret": .25, "own": .20, "agi": .15, "spd": .10}
O = {"1": [90, 40, 30, 35, 95], "2": [95, 55, 45, 30, 80], "3": [100, 95, 85, 90, 75]}
keys = list(W)
base_w = [W[k] for k in keys]
names = {"cap": "Meets $25M cap", "ret": "Growth return", "own": "GM ownership", "agi": "Agility", "spd": "Speed"}


def score(wv):
    return {k: sum(a * b for a, b in zip(wv, v)) for k, v in O.items()}


rows = [["Weight changed", "Option 1", "Option 2", "Option 3", "Order"]]
b = score(base_w)
rows.append([[("Base weights", True)], f"{b['1']:.1f}", f"{b['2']:.1f}", f"{b['3']:.1f}", "3 > 2 > 1"])
for i, k in enumerate(keys):
    for dlt in (+.10, -.10):
        nw = base_w[:]
        nw[i] = max(0, nw[i] + dlt)
        rest = 1 - nw[i]
        tot = sum(base_w[j] for j in range(5) if j != i)
        for j in range(5):
            if j != i:
                nw[j] = base_w[j] / tot * rest
        sc = score(nw)
        order = " > ".join(sorted(sc, key=sc.get, reverse=True))
        rows.append([f"{names[k]} {'+' if dlt > 0 else '-'}10 pts", f"{sc['1']:.1f}", f"{sc['2']:.1f}", f"{sc['3']:.1f}", order])
table(s, 0.67, 1.45, [3.4, 1.7, 1.7, 1.7, 1.9], rows, size=10.5, header_size=10.5, row_min=0.38, max_bottom=6.75)
footer(s, "Base weights: cap 30%, growth return 25%, GM ownership 20%, agility 15%, speed 10%. The changed weight moves 10 points; "
          "the others are rebalanced pro rata. Smallest gap between options 3 and 2: 24.2 points. Scores are judgment (A9).",
       y=6.8, h=0.45)
notes(s, "Backup.")

# ================================================================== D. Sources
s = new_slide(title="D")
header(s, "BACKUP D", "Every number is a sourced fact, a labeled assumption or context")
rows = [["Topic", "Source and class", "Data point", "Where it is used"],
        ["Nu scale and efficiency", "Nu Q2'26 release and 6-K; Form 20-F  |  FACT", "138.9M customers; ARPAC $17.1; cost to serve $1.0 ($0.8 a year earlier); efficiency ratio 19.5% (17.6% in Q1; ~20% FY26 per the call); ROE 33%; net income $1.06B; opex $806.2M; net revenue ~$4.1B; low-and-grow underwriting.", "Slides 3, 4, 7; backup I"],
        ["Corporate events", "Nu announcements; OCC; Nu IR  |  FACT", "Mexico bank authorization 10 Jul 2026 (operating from 6 Aug), $4.2B plan to 2030; OCC conditional approval 29 Jan 2026 (capital within 12 months); US launch via Lead Bank 10 Sep 2026; buyback up to $1.0B (4 Jun 2026); CFO since 13 Jul 2026; Q3'26 results 12 Nov; Investor Day 8 Dec 2026.", "Slides 7, 9, 15; backup B"],
        ["Cohorts", "Q1'26 earnings presentation  |  FACT", "$30.8 monthly ARPAC at 96 months since origination vs $15.9 total.", "Slide 7; backup E"],
        ["HQ tax rules", "OECD GloBE Model Rules; Form 20-F  |  FACT (rule)", "EUR 750M threshold; 15% minimum; 2026 transitional carve-out 9.4% of payroll and 7.4% of tangible assets (Art. 5.3, 9.2); Nu Holdings is Cayman-incorporated.", "Slide 10"],
        ["Context only", "Hackett, McKinsey, Gartner, HBR, NACD, Deloitte, Baker McKenzie; three-horizons framing (Baghai, Coley and White)  |  CONTEXT", "Central hubs shorten planning cycles; active reallocators earn higher returns; fair process wins acceptance; Board books at T-7.", "Footers and talk only"],
        ["AI in operations", "Nu Q2'26 release; Nu newsroom (Hyperplane, Jun 2024)  |  FACT. Gartner; PMI; Science; HBS/BCG  |  CONTEXT", "AI agents handle more than 60% of support chats in Brazil; AI-first strategy. Studies and surveys on AI in project reporting, writing and finance (backup H).", "Slides 6, 9, 12; backup H"]]
bottom = table(s, 0.67, 1.45, [2.0, 2.9, 5.1, 2.0], rows, size=10, header_size=10.5, row_min=0.36, max_bottom=6.6)
footer(s, "Research assisted by AI tools; every figure, calculation and recommendation was checked by me.", y=max(bottom + 0.15, 6.6), h=0.35)
notes(s, "Backup.")

# ================================================================== E. Full-year split
s = new_slide(title="E")
header(s, "BACKUP E", "Full-year split: core markets protect the year, new bets earn the reserve")
card(s, 0.67, 1.8, 5.8, 3.4, "CORE MARKETS  //  BRAZIL, MEXICO, COLOMBIA", "$16.5M in the cap, plus up to $5.0M from the reserve",
     ["70% of the first release: $10.5M on 1 January.",
      "Mature cohorts reach about $31 of monthly ARPAC at eight years ($30.8 at 96 months, Q1'26) against $17.1 for the group, at 33% ROE."],
     body_size=12.5)
card(s, 6.7, 1.8, 5.8, 3.4, "NEW BETS  //  NEW VERTICALS AND MARKETS", "$8.5M in the cap, plus up to $10.0M from the reserve",
     ["30% of the first release: $4.5M on 1 January.",
      "Judged on activation, retention and unit cost until a cohort can show payback. First claim on up to $14M across the "
      "second release and the reserve, if it opens."], body_size=12.5)
band(s, 0.67, 5.45, 11.8, 0.9, ["A new bet judged on mature-market profit in its first quarter always loses. "
                                "The horizon split is a frame, not evidence."], size=13, dark=False)
notes(s, "Backup.")

# ================================================================== F. Day-by-day timeline
s = new_slide(title="F")
header(s, "BACKUP F", "One timeline: tenure days, Board countdown and HQ gates")
rows = [["Tenure day", "Date", "Board countdown / HQ gate", "What happens"],
        ["Day 60", "Tue 20 Oct", "T-28  |  HQ D+0", "Board outline out; HQ charter signed; tax gate launched"],
        ["Day 67", "Tue 27 Oct", "T-21", "One-to-ones with each executive"],
        ["Day 72", "Early Nov", "T-16", "Capital options lock with CEO, CFO and MTeam"],
        ["Day 78-81", "Early Nov", "T-10 to T-7", "Inputs freeze; chairs briefed (T-9); book on portal (T-7)"],
        ["Day 88", "Tue 17 Nov (assumed)", "T-0", "Board approves FY27 plan; first-release allocation published"],
        ["Day 90", "Thu 19 Nov", "HQ D+30", "Gate 1: tax and residence opinion; two sites to term sheet"],
        ["Day 95", "Late Nov", "", "Hub in place; finance partners embedded"],
        ["Day 105", "Early Dec", "", "Squad clinics running; $12M case track under way"],
        ["Day 120", "Mid-Dec", "HQ D+60", "Gate 2: grant terms (24-month grace); one site dropped"],
        ["Day 133", "Fri 1 Jan 2027", "", "FY27 starts; $15M first release"],
        ["Day 150", "Mid-Jan", "HQ D+90", "Gate 3: officers and D&O; 90-day report to the CFO"],
        ["Day 180", "Mid-Feb", "HQ D+120", "Gate 4: lease"],
        ["Day 222", "Wed 31 Mar", "", "Q1 close; second release ($10M) decided on Q1 data in early April"]]
table(s, 0.67, 1.5, [1.5, 2.3, 2.7, 5.5], rows, size=10.5, header_size=10.5, row_min=0.36, max_bottom=7.0)
notes(s, "Backup.")

# ================================================================== G. Stakeholder map
s = new_slide(title="G")
header(s, "BACKUP G", "Stakeholder map: who decides, who is consulted, and when they hear it")
rows = [["Who", "Role in these decisions", "When they hear it"],
        ["CEO", "Decides the capital ranking", "Options memo before T-16; CEO's office briefed before any GM lock"],
        ["CFO (since 13 Jul 2026)", "Recommends the ranking; decides releases and the H2 reserve", "Weekly; every pack before it moves"],
        ["Board chair and committee chairs", "Approve the capital plan; oversee entities", "Pre-read at T-9; book at T-7"],
        ["Corporate Secretary", "Owns the Board book and statutory calendar", "From today (T-28)"],
        ["The two MTeam members", "Hold the opposing capital views", "One-to-ones at T-21; joint lock at T-16"],
        ["Local CFOs (Mexico, Colombia); Brazil CFO (new role)", "Own local decision rights; inputs to the model", "Squad clinics; decision rights unchanged for 90 days"],
        ["Risk", "Confirms credit conditions; can block the H2 reserve", "Quarterly release review; H1 close"],
        ["IR", "Sets how the plan and reserve are disclosed", "Before Investor Day (8 Dec)"],
        ["Tax, Legal, Compliance, People, Workplace", "Own HQ gates 1-4", "Weekly HQ pod; 48-hour escalation"],
        ["Counterparties (local authority, incentive agency, landlord)", "Across the table", "Only after gate 1, through Legal; never before the CEO and CFO see the scorecard"],
        ["GMs", "Bring the cases; co-sign them", "Squad clinics weekly; one-to-ones every two weeks"]]
table(s, 0.67, 1.5, [3.2, 4.4, 4.4], rows, size=10.5, header_size=10.5, row_min=0.38, max_bottom=7.0)
notes(s, "Backup.")

# ================================================================== H. AI in the operating model
s = new_slide(title="H")
header(s, "BACKUP H", "Two AI agents take the busywork; people own every number and decision")
rows = [["", "PMO agent", "Review agent"],
        ["Day to day", "Every Monday it reads status where the work happens (tracker, shared drive, team channels) for the HQ gates, the "
         "release plan and the Board countdown. It updates the gate tracker and decision log, flags dates at risk and drafts the pod note.",
         "When a GM submits a case, it checks the template and the four indicators, ties each figure to the data platform and flags "
         "gaps before the clinic. For the Board, it drafts the book from the locked tables and drafts answers to directors' questions."],
        ["What it replaces", "Status templates and the weekly chase. Owners correct a draft instead of filling in a form.",
         "Manual tie-outs and blank-page first drafts. Analyst time moves to judgment."],
        ["Why it is good practice", "Gartner (2019): data collection, tracking and reporting are the project-management tasks AI takes "
         "over; 80% of today's tasks gone by 2030. PMI (2023): reporting is project managers' top AI use (34%).",
         "Science (Noy and Zhang, 2023): writing tasks 40% faster, quality up 18%. HBS/BCG (2023): inside AI's frontier, 12.2% more "
         "tasks and 40% higher quality; outside it, 19 points less accurate. Gartner (2024): 58% of finance teams use AI."],
        ["Guardrails", "Reads only; never changes a date or a decision. Each owner confirms their line. Escalation to the CFO stays "
         "human, within 48 hours.",
         "Nu-approved tools inside Nu's environment (A11). A named person signs every figure and answer. A wrong number that reaches "
         "the CFO sends that step back to manual until it is fixed."],
        ["How I measure it", "Hours spent on status (baseline now); gate dates hit; no status templates by late November.",
         "Cycle time under 40 days; errors caught before the clinic; directors' answers out by T-3."]]
table(s, 0.67, 1.45, [1.9, 5.05, 5.05], rows, size=11, header_size=11.5, row_min=0.5, max_bottom=6.55)
footer(s, "Nu context: AI agents handle more than 60% of support conversations in Brazil (Q2'26 release); AI-first strategy "
          "(Hyperplane acquisition, Jun 2024). The studies are context for the method, not a forecast of Nu's savings.", y=6.6, h=0.4)
notes(s, "Backup.")

# ================================================================== I. Scenarios and stop rules
s = new_slide(title="I")
header(s, "BACKUP I", "The downside is handled by rules: miss Q1 and the second release waits")
rows = [["", "Market slows", "Base plan", "H1 beats plan"],
        ["Funded spend vs the $25M cap", "$15.0M (second release held at Q1 close)", "$25.0M ($15M + $10M)", "$25.0M"],
        ["H2 reserve", "$0", "$0 (not assumed)", "Up to $15.0M from H1 profit above plan, inside the guide in force"],
        ["Stop rules", "A unit below plan on two of four indicators: its second-release share is held",
         "Efficiency ratio above the guide in force at Q1 close: the second release is held for every unit",
         "Credit cost or NPL 90+ above plan: the reserve stays shut; Risk can block it"],
        ["Top risks and owners", "GMs game the indicators: finance partners own the read-out",
         "Board book slips: the outline locked today decides what goes to the appendix",
         "HQ: tax residence, clawback, officer approval, site politics, key-person risk (backup J)"],
        ["What I report in 90 days", "First monthly indicator read-outs per unit", "Board delivered; capital plan approved; cycle-time and forecast baselines",
         "HQ gates 1 and 2 cleared; gate 3 pack ready"]]
bottom = table(s, 0.67, 1.75, [2.4, 3.1, 3.2, 3.1], rows, size=11, header_size=11.5, row_min=0.55)
footer(s, "Margins are protected without a mid-year freeze. No value figures are claimed beyond the spend lines; before HQ gate 1 "
          "they would be guesses.", y=max(bottom + 0.2, 6.4), h=0.45)
notes(s, "Backup.")

# ================================================================== J. HQ site scorecard and risk register
s = new_slide(title="J")
header(s, "BACKUP J", "HQ site scorecard and risk register: hard tests first, then weights")
add_text(s, 0.67, 1.6, 5.9, 0.75, [[("Hard tests first, pass or fail: ", True), ("license, tax residence and taxable presence, "
         "fit-and-proper officers, real payroll and assets. Only sites that pass are scored.", False)]], size=11, color=INK, where="hard tests")
rows = [["Criterion", "Weight", "Evidence and owner"],
        ["Talent pool and wage cost", "30", "Hiring data and wage benchmarks (People)"],
        ["Incentive value net of clawback", "20", "Draft terms with the clawback modeled (Legal, Tax)"],
        ["Campus cost, flexibility and timing", "20", "Lease options, break clauses, fit-out time (Workplace)"],
        ["Modeled tax rate under Pillar Two, including any domestic top-up", "15", "The gate 1 model (Tax)"],
        ["Political and reputational risk", "15", "Local stakeholder map (Legal, Comms)"]]
table(s, 0.67, 2.45, [2.75, 0.75, 2.4], rows, size=10.5, header_size=11, row_min=0.42, max_bottom=6.5)
rows = [["Risk", "Mitigation", "Owner"],
        ["Tax residence or a taxable presence moves", "Opinion at gate 1; meeting patterns unchanged; no early signing", "Tax"],
        ["Grant clawback", "Proportional scaling and 24 months' grace; walk away from cliff-edge terms", "Legal"],
        ["Officer approval runs late", "Files at gate 1; no activation before approval; the lease long-stop moves with it", "Legal, People"],
        ["Site politics and disclosure", "No public contact before the scorecard; one spokesperson (Comms); IR decides what is disclosed and when", "Comms, IR"],
        ["Dependence on me", "A named deputy, the decision log, gate packs the CFO can read without me", "Me"]]
table(s, 6.8, 1.6, [1.95, 2.85, 0.9], rows, size=10.5, header_size=11, row_min=0.42, max_bottom=6.5)
footer(s, "Weights and mitigations are my proposals (A13). No savings figure until a site is chosen.", y=6.7, h=0.35)
notes(s, "Backup.")

# title collision guard: nothing may start between the title's top and its predicted last line
for i, sl in enumerate(prs.slides, 1):
    tb = CUR.get("title_bottom", {}).get(i)
    if tb is None:
        continue
    for shp in sl.shapes:
        top = shp.top / 914400
        if 0.9 < top < tb + 0.05:
            WARN.append(f"slide {i}: shape '{shp.name}' at {top:.2f}in collides with title ending {tb:.2f}in")

number_all()
prs.save(DST)
pptx_safe.make_powerpoint_safe(DST)
print("saved", DST, "slides:", len(prs.slides))
print("FIT WARNINGS:" if WARN else "fit check: no overflow predicted")
for w in WARN:
    print("  ", w)
